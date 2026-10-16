import io
import openpyxl
from fastapi.testclient import TestClient
from app.main import app

def client():
    c = TestClient(app); c.__enter__(); c.post("/auth/dev-login", json={"email": "manager@local"}); return c

def test_sample_sheet_needs_sign_in():
    c = TestClient(app); c.__enter__()
    assert c.get("/api/demo/sample-sheet").status_code == 401

def test_sample_sheet_is_a_valid_import_file():
    c = client(); r = c.get("/api/demo/sample-sheet")
    assert r.status_code == 200 and "spreadsheetml" in r.headers["content-type"]
    ws = openpyxl.load_workbook(io.BytesIO(r.content)).active
    assert [x.value for x in ws[1]] == ["Team Member", "Current Experience", "Experience Bracket", "Role per MSA", "Rate Card", "Genus Cost", "Margin"]
    assert ws.max_row == 9                                            # header + 8 invented people
    tid = c.post("/api/teams", json={"name": "Guided demo import check", "calendar_id": c.get("/api/calendars").json()[0]["id"]}).json()["id"]
    imp = c.post(f"/api/team/import?team_id={tid}", files={"file": ("sample_team_sheet.xlsx", r.content, "application/octet-stream")})
    assert imp.status_code == 200 and len(imp.json()) == 8
    assert all(not any("not in Genus table" in w for w in p["warnings"]) for p in imp.json())   # every sample role exists in the Genus table

def test_guided_demo_assets_are_served():
    c = TestClient(app); c.__enter__()
    assert c.get("/static/js/demo.js").status_code == 200

def test_tour_step_ids_are_unique():
    import re
    from pathlib import Path
    js = (Path(__file__).resolve().parents[2] / "frontend" / "js" / "demo.js").read_text()
    ids = re.findall(r"^\{id:'([a-z-]+)'", js, re.M)
    assert len(ids) > 20 and len(ids) == len(set(ids))             # progress is saved by id, so a repeated id would loop the tour
